# -*- coding: utf-8 -*-
"""Verify the presentation, rehearsal documents and cited technical facts.

Run from the repository root after installing the presentation toolchain:

    python3 presentation/tools/verify_rehearsal.py

The check intentionally fails on stale wording that previously appeared in the
deck, so those errors cannot quietly return.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
PRESENTATION = ROOT / "presentation"
sys.path.insert(0, str(ROOT))

import din6892_methods as din  # noqa: E402
import fva600_matlab as matlab  # noqa: E402
import fva600_postprocess as post  # noqa: E402
import keyjoint_core as core  # noqa: E402
import model_builder_i18n as i18n  # noqa: E402
import project_workspace as workspace  # noqa: E402

checks = []


def check(label, condition, detail=""):
    checks.append((label, bool(condition), detail))
    print("%-64s %s%s" % (
        label,
        "OK" if condition else "FAIL",
        ("  " + str(detail)) if detail else "",
    ))


def close(actual, expected, tolerance=1.0e-9):
    return abs(float(actual) - float(expected)) <= tolerance


def deck_text(path):
    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        slides.append("\n".join(
            shape.text_frame.text for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()))
    return prs, slides


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def parse_slide_route(document):
    """Return the normalized slide numbers from the canonical route sentence."""
    match = re.search(r"present slides ([^;]+);", document)
    if not match:
        raise ValueError("Canonical 20-minute route sentence was not found")
    value = match.group(1).replace(" and ", ", ")
    numbers = []
    for token in (item.strip() for item in value.split(",")):
        if not token:
            continue
        if "–" in token:
            start, end = (int(part.strip()) for part in token.split("–", 1))
            numbers.extend(range(start, end + 1))
        else:
            numbers.append(int(token))
    return tuple(numbers)


def main():
    print("=== DOCUMENT STRUCTURE ===")
    english, english_slides = deck_text(
        PRESENTATION / "ModelBuilder_4.2_Presentation.pptx")
    spanish, spanish_slides = deck_text(
        PRESENTATION / "ModelBuilder_4.2_Presentacion.pptx")
    check("English deck has 39 slides", len(english.slides) == 39,
          len(english.slides))
    check("Spanish deck has 39 slides", len(spanish.slides) == 39,
          len(spanish.slides))

    script = (PRESENTATION / "PRESENTATION_SCRIPT_EN_ES.md").read_text(
        encoding="utf-8")
    questions = (PRESENTATION / "SUPERVISOR_QA_EN_ES.md").read_text(
        encoding="utf-8")
    cheat = (PRESENTATION / "REHEARSAL_CHEAT_SHEET_EN_ES.md").read_text(
        encoding="utf-8")
    script_numbers = [int(value) for value in re.findall(
        r"^## Slide (\d+)\b", script, flags=re.MULTILINE)]
    question_numbers = [int(value) for value in re.findall(
        r"^## Q(\d+)\.", questions, flags=re.MULTILINE)]
    check("Speaker script covers slides 1..39",
          script_numbers == list(range(1, 40)), script_numbers[-3:])
    check("Q&A bank covers questions 1..69",
          question_numbers == list(range(1, 70)), question_numbers[-3:])
    check("Cheat sheet contains opening, closing and pronunciation",
          all(term in cheat for term in (
              "# 2. Opening and closing", "# 7. Pronunciation",
              "# 10. Final pre-meeting checklist")))

    print("\n=== STALE OR UNSAFE WORDING ===")
    text_sources = "\n".join((
        script, questions, cheat,
        (PRESENTATION / "tools" / "make_deck.py").read_text(encoding="utf-8"),
        (PRESENTATION / "tools" / "make_diagrams.py").read_text(encoding="utf-8"),
    ))
    unsafe_current_claims = (
        "Every number in this deck is reproducible",
        "Cada número de esta presentación es reproducible",
        "K_λ, K_R, f_H and f_S stay at 1.0",
        "K_λ, K_R, f_H y f_S siguen en 1,0",
        "matching mesh over 20 cycles",
        "malla coincidente de 20 ciclos",
        "D/32 on the shaft, b/20 on the key and d_a/80 on the hub",
        "D/32 en el eje, b/20 en la chaveta y d_a/80 en el cubo",
        "the real allowable torque drops another 46",
        "el par admisible real baja otro 46",
        "BUILD_RESULT, CAE, screenshots and report all exist",
        "BUILD_RESULT, CAE, capturas e informe",
    )
    for phrase in unsafe_current_claims:
        check("Stale phrase absent: %s" % phrase[:42], phrase not in text_sources)
    check("Slide 34 source says eight mesh templates",
          "eight mesh templates" in english_slides[33] or
          "eight mesh templates" in (PRESENTATION / "tools" / "make_deck.py").read_text(
              encoding="utf-8"))
    check("Slide 34 does not claim four packaged presets",
          "All four presets" not in english_slides[33])
    check("Closing qualifies historical evidence",
          "source evidence" in english_slides[38])
    check("Regular-model slide does not infer matching from image",
          "not inferred from the image" in english_slides[16])
    check("Deck states artifact requirements conditionally",
          "requested or required artifacts" in english_slides[33])
    gui_source = (ROOT / "model_builder_gui.py").read_text(encoding="utf-8")
    check("Artifact contract conditions CAE and screenshots on settings",
          'if params.get("save_cae")' in gui_source and
          'if params.get("make_preview")' in gui_source)

    slide_seconds = [int(value) for value in re.findall(
        r"^\*\*Time:\*\* (\d+) seconds", script, flags=re.MULTILINE)]
    script_route = parse_slide_route(script)
    cheat_route = parse_slide_route(cheat)
    check("Speaker script provides one time for every slide",
          len(slide_seconds) == 39, len(slide_seconds))
    check("Full script timings sum to 31 minutes",
          sum(slide_seconds) == 31 * 60, sum(slide_seconds))
    check("Script and cheat sheet parse to the same 20-minute route",
          script_route == cheat_route, script_route)
    route_seconds = sum(slide_seconds[number - 1] for number in script_route)
    check("Canonical abbreviated route is no longer than 20 minutes",
          route_seconds <= 20 * 60, route_seconds)
    check("Cheat-sheet displayed timing total agrees with parsed total",
          "| **Total** | **39** | **%d min** |" %
          (sum(slide_seconds) // 60) in cheat)
    readme = (PRESENTATION / "README.md").read_text(encoding="utf-8")
    check("README lists and shows how to run verify_rehearsal.py",
          "tools/verify_rehearsal.py" in readme and
          "python3 presentation/tools/verify_rehearsal.py" in readme)

    print("\n=== CATALOGUE AND SCOPE COUNTS ===")
    catalog = i18n.validate_catalogs(str(ROOT / "i18n"))
    check("DIN 6885 has 26 rows", len(core.DIN6885) == 26,
          len(core.DIN6885))
    check("DIN 6885 has 34 standard lengths",
          len(core.DIN6885_LENGTHS) == 34, len(core.DIN6885_LENGTHS))
    check("DIN 6885 has 7 radius bands", len(core.DIN6885_RADII) == 7,
          len(core.DIN6885_RADII))
    check("Current code has 8 mesh templates", len(core.MESH_TEMPLATES) == 8,
          len(core.MESH_TEMPLATES))
    check("Current code has 15 presets", len(core.PRESETS) == 15,
          len(core.PRESETS))
    check("Method catalogue has four identifiers", len(din.METHOD_IDS) == 4,
          din.METHOD_IDS)
    check("FVA catalogue has VB1..VB8", sorted(din.FVA_VARIANTS) == [
        "VB%d" % index for index in range(1, 9)])
    check("i18n has es/en/de and 650 keys",
          tuple(catalog["locales"]) == ("es", "en", "de") and
          len(catalog["keys"]) == 650,
          "%s / %d" % (catalog["locales"], len(catalog["keys"])))
    check("Workspace has the seven cited directories",
          tuple(workspace.DIRECTORY_NAMES) == (
              "input", "generated", "model", "jobs", "screenshots",
              "reports", "logs"), workspace.DIRECTORY_NAMES)

    print("\n=== DEFAULT DIN VALUES AND CORRECTIONS ===")
    params = core.default_params()
    derived = core.derive(copy.deepcopy(params))
    din_record = derived["din6892"]
    calculation = din_record["calculation"]
    consistency = din_record["companions"]["consistency"]["torques_Nm"]
    effective = din_record["derived"]["effective_bearing_depth"]
    hub_effective = din_record["derived"]["hub_effective_bearing_depth"]
    row = core.din6885_row(40.0)
    radii = core.din6885_radii(12.0)
    check("D40 selects band 38 < d1 <= 44",
          (row["D_from"], row["D_upto"]) == (38.0, 44.0))
    check("D40 key row is 12 x 8, t1=5.0, t2=3.3",
          (row["b"], row["h"], row["t1"], row["t2"]) ==
          (12.0, 8.0, 5.0, 3.3))
    check("D40 radii are r1 0.25..0.40 and r2 0.16..0.25",
          (radii["r1_min"], radii["r1_max"],
           radii["r2_min"], radii["r2_max"]) ==
          (0.25, 0.40, 0.16, 0.25))
    check("Current Method B is 999.1827188 N.m",
          close(consistency["B_DIN_CURRENT"], 999.1827188023801))
    check("Method C is 572.508 N.m",
          close(consistency["C_PRELIMINARY"], 572.508))
    check("FVA B diagnostic is 2464.866485 N.m",
          close(consistency["B_FVA_2025"], 2464.86648545341))
    check("Current Method B is hub-governed",
          calculation["governing_component"] == "hub")
    check("t1tr literal, t2tr hub and t1tr geometry match cited values",
          close(effective["t1tr_mm"], 5.141491287186004) and
          close(hub_effective["t2tr_mm"], 3.1414912871860032) and
          close(effective["t1tr_geometric_mm"], 2.7585087128139962))
    check("Equation-9 spread is 46.3481 percent",
          close(effective["relative_difference_percent"],
                46.348081544182506))
    check("Hub area was over-credited by 63.6640 percent",
          close((effective["t1tr_mm"] / hub_effective["t2tr_mm"] - 1.0) * 100,
                63.664031415840874))
    fw = din.load_reversal_factor(1.0e4)["f_W"]
    alternating_inputs = dict(
        (key, value) for key, value in
        din_record["derived"]["calculation_inputs"].items()
        if key != "N_W")
    alternating_inputs.update({
        "f_S": calculation["f_S"], "f_H": calculation["f_H"],
        "N_W": 1.0e4,
    })
    alternating = din.method_b_current(alternating_inputs)[
        "torque_allowable_Nm"]
    check("f_W at 10,000 reversals is 0.7962143411",
          close(fw, 0.7962143411069944))
    check("Fully alternating torque is 795.5636101 N.m",
          close(alternating, 795.5636100967324))
    reconstructed = (consistency["B_DIN_CURRENT"] *
                     effective["t1tr_mm"] / hub_effective["t2tr_mm"] * fw)
    hub_fix_only = consistency["B_DIN_CURRENT"] * fw
    check("Pre-correction 1302.05 N.m reconstructs from both defects",
          close(reconstructed, 1302.0515, tolerance=5.0e-4), reconstructed)
    check("Hub-depth fix alone with old f_W gives 795.5636101 N.m",
          close(hub_fix_only, 795.5636100967324))
    check("Script identifies 1302.05 -> 999.18 as combined net effect",
          "combined net effect of both corrections" in script)

    policy_results = {}
    shaft_results = {}
    for policy in ("EQUATION_9", "GEOMETRIC", "CONSERVATIVE"):
        probe = core.default_params()
        probe["din6892"]["bearing_depth_policy"] = policy
        record = core.derive(probe)["din6892"]["calculation"]
        policy_results[policy] = (
            record["torque_allowable_Nm"], record["governing_component"])
        shaft_results[policy] = next(
            row["torque_allowable_Nm"] for row in record["components"]
            if row["name"] == "shaft")
    check("All three depth policies remain hub-governed at 999.1827188 N.m",
          all(close(value[0], 999.1827188023801) and value[1] == "hub"
              for value in policy_results.values()), policy_results)
    check("Geometric policy reduces shaft component to 1027.4782913 N.m",
          close(shaft_results["GEOMETRIC"], 1027.4782913141062),
          shaft_results["GEOMETRIC"])
    check("Deck does not claim equation 9 lowers governing D40 torque",
          "hub still governs at 999.18 N·m" in english_slides[37])

    print("\n=== DEFAULT VALIDATION, SEEDS AND CONFIGURATIONS ===")
    issues = core.validate(copy.deepcopy(params), derived)
    check("Default has 0 errors and 3 warnings",
          len(core.errors(issues)) == 0 and len(core.warnings_(issues)) == 3,
          [item["code"] for item in core.warnings_(issues)])
    check("Default seed rules are D/16, b/20 and d_a/40",
          derived["seed_rule_shaft"] == "D/16" and
          derived["seed_rule_hub"] == "d_a/40" and
          close(derived["seed_key"], derived["b"] / 20.0))
    packaged_configurations = (
        ROOT / "params_default.json",
        ROOT / "params_fva600_research_presolve_1lw.json",
        ROOT / "params_fva600_method_a_20lw_nojob.json",
    )
    packaged_ok = True
    for path in packaged_configurations:
        payload = core.normalize_params(json.loads(path.read_text(encoding="utf-8")))
        packaged_ok = packaged_ok and not core.errors(
            core.validate(payload, core.derive(payload)))
    ten_cycle_path = ROOT / "params_fva600_method_a_10lw_solve.json"
    ten_cycle = core.normalize_params(json.loads(
        ten_cycle_path.read_text(encoding="utf-8")))
    spec_text = (ROOT / "ModelBuilder.spec").read_text(encoding="utf-8")
    check("Three packaged configurations validate under schema 6", packaged_ok)
    check("10-cycle repository setup validates but is not packaged",
          not core.errors(core.validate(ten_cycle, core.derive(ten_cycle))) and
          ten_cycle_path.name not in spec_text)

    print("\n=== METHOD A AND MATLAB CONTRACT ===")
    check("Method A CSV columns are x_mm, z_mm, opening_um",
          post.REQUIRED_COLUMNS == ("x_mm", "z_mm", "opening_um"))
    check("Method A coverage tolerance is 2 percent",
          close(post.COVERAGE_TOLERANCE, 0.02))
    check("MATLAB modes are OFF, EXPORT, RUN",
          matlab.MODES == ("OFF", "EXPORT", "RUN"))
    check("MATLAB tolerances are 1e-10 absolute and 1e-9 relative",
          matlab.DEFAULT_ABS_TOLERANCE == 1.0e-10 and
          matlab.DEFAULT_REL_TOLERANCE == 1.0e-9)
    matlab_source = (ROOT / "fva600_matlab.py").read_text(encoding="utf-8")
    check("MATLAB bundle manifest is exactly manifest.json",
          'os.path.join(bundle_dir, "manifest.json")' in matlab_source)
    check("MATLAB source states Python is authoritative",
          '"python_authoritative": True' in matlab_source and
          "result never replaces the Python value used by Model Builder" in
          matlab_source)

    print("\n=== BUILD AND PUBLISHED-EXE CLAIMS ===")
    import make_release  # noqa: E402
    check("Release build fingerprints 31 inputs",
          len(make_release.BUILD_INPUT_FILES) == 31)
    batch = (ROOT / "make_exe.bat").read_text(encoding="utf-8")
    check("Build pins PyInstaller 6.22.2 and hooks 2026.7",
          "pyinstaller==6.22.2" in batch and
          "pyinstaller-hooks-contrib==2026.7" in batch)
    manifest = json.loads((
        ROOT / "release" / "ModelBuilder_4.2" / "BUILD_MANIFEST.json"
    ).read_text(encoding="utf-8"))
    records = {
        item["path"].split("parametric_builder/", 1)[-1]: item
        for item in manifest["build_provenance"]["source_inputs"]
        if item["path"].startswith("parametric_builder/")
    }
    changed = sum(
        1 for relative, item in records.items()
        if (ROOT / relative).is_file() and sha256(ROOT / relative) != item["sha256"])
    check("Published EXE manifest has 31 source inputs",
          len(manifest["build_provenance"]["source_inputs"]) == 31)
    check("17 parametric-builder inputs changed since published EXE",
          changed == 17, changed)

    failed = [label for label, passed, _detail in checks if not passed]
    print("\n=== RESULT ===")
    if failed:
        print("%d check(s) failed:" % len(failed))
        for label in failed:
            print("  - " + label)
        return 1
    print("All %d checks passed." % len(checks))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

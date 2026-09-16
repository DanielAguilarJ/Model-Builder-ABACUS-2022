# -*- coding: utf-8 -*-
"""Rasterise the saved .pptx so the slides can be inspected visually.

This reads the real file back with python-pptx and draws it with PIL using
Liberation Sans, which is metric-compatible with Arial, so text wrapping and
overflow match what PowerPoint will do closely enough to catch layout bugs.

Usage:  python3 render_preview.py [slide numbers ...]
"""
from __future__ import annotations

import io
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu

DECK = "/projects/sandbox/Model-Builder-ABACUS-2022/presentation/ModelBuilder_4.2_Presentacion.pptx"
OUT = "/projects/sandbox/assets/slides_png"
PPI = 120.0
EMU_PER_INCH = 914400.0

FONTS = {
    ("Arial", False): "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
    ("Arial", True): "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf",
    ("Consolas", False): "/usr/share/fonts/liberation-mono/LiberationMono-Regular.ttf",
    ("Consolas", True): "/usr/share/fonts/liberation-mono/LiberationMono-Bold.ttf",
}
_CACHE = {}


def font_for(name, bold, size_px):
    path = FONTS.get((name, bold)) or FONTS[("Arial", bold)]
    if not os.path.isfile(path):
        for candidate in ("/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
                          "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
            if os.path.isfile(candidate):
                path = candidate
                break
    key = (path, int(size_px))
    if key not in _CACHE:
        _CACHE[key] = ImageFont.truetype(path, max(int(round(size_px)), 6))
    return _CACHE[key]


def px(emu):
    return int(round(emu / EMU_PER_INCH * PPI))


def rgb(color, default=(0, 0, 0)):
    try:
        value = color.rgb
        return (value[0], value[1], value[2])
    except Exception:
        return default


def wrap(text, font, max_width, draw):
    lines = []
    for hard_line in text.split("\n"):
        if not hard_line:
            lines.append("")
            continue
        words = hard_line.split(" ")
        current = ""
        for word in words:
            candidate = (current + " " + word).strip()
            width = draw.textlength(candidate, font=font)
            if width > max_width and current:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def draw_text_frame(draw, shape, offset=(0, 0)):
    frame = shape.text_frame
    left = px(shape.left) + offset[0]
    top = px(shape.top) + offset[1]
    width = px(shape.width)
    height = px(shape.height)
    lines = []
    total = 0
    for paragraph in frame.paragraphs:
        runs = paragraph.runs
        if not runs:
            continue
        run = runs[0]
        size_pt = run.font.size.pt if run.font.size else 18.0
        size_px = size_pt * PPI / 72.0
        font = font_for(run.font.name or "Arial", bool(run.font.bold), size_px)
        color = rgb(run.font.color, (0, 0, 0))
        spacing = paragraph.line_spacing or 1.0
        before = paragraph.space_before.pt * PPI / 72.0 if paragraph.space_before else 0.0
        text_value = "".join(r.text for r in runs)
        wrapped = wrap(text_value, font, width, draw)
        line_height = size_px * 1.2 * float(spacing)
        lines.append((wrapped, font, color, line_height, before,
                      paragraph.alignment))
        total += before + line_height * len(wrapped)

    anchor = frame.vertical_anchor
    y = top
    if anchor == MSO_ANCHOR.MIDDLE:
        y = top + (height - total) / 2.0
    overflow = total - height
    for wrapped, font, color, line_height, before, alignment in lines:
        y += before
        for line in wrapped:
            text_width = draw.textlength(line, font=font)
            if alignment == PP_ALIGN.CENTER:
                x = left + (width - text_width) / 2.0
            elif alignment == PP_ALIGN.RIGHT:
                x = left + width - text_width
            else:
                x = left
            ascent, _descent = font.getmetrics()
            draw.text((x, y + (line_height - font.size * 1.2) / 2.0), line,
                      font=font, fill=color)
            y += line_height
    return overflow


def rounded(draw, box, radius, fill):
    draw.rounded_rectangle(box, radius=radius, fill=fill)


def render(slide, index):
    width = px(Emu(12192000))
    canvas = Image.new("RGB", (1600, 900), (255, 255, 255))
    try:
        background = rgb(slide.background.fill.fore_color, (255, 255, 255))
    except Exception:
        background = (255, 255, 255)
    canvas.paste(Image.new("RGB", canvas.size, background), (0, 0))
    draw = ImageDraw.Draw(canvas)
    warnings = []

    for shape in slide.shapes:
        if shape.shape_type == 13 or shape.__class__.__name__ == "Picture":
            try:
                image = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
            except Exception as exc:
                warnings.append("picture failed: %s" % exc)
                continue
            target = (max(px(shape.width), 1), max(px(shape.height), 1))
            image = image.resize(target, Image.LANCZOS)
            canvas.paste(image, (px(shape.left), px(shape.top)), image)
            continue
        if shape.has_text_frame and shape.text_frame.text.strip() and \
                shape.shape_type is not None and \
                shape.__class__.__name__ == "Shape" and shape.fill.type == 1:
            # filled autoshape carrying text (a chip)
            fill = rgb(shape.fill.fore_color, (240, 240, 240))
            box = [px(shape.left), px(shape.top),
                   px(shape.left) + px(shape.width),
                   px(shape.top) + px(shape.height)]
            rounded(draw, box, radius=min(px(shape.height) // 2, 20), fill=fill)
            draw_text_frame(draw, shape)
            continue
        if not shape.has_text_frame:
            continue
        if shape.__class__.__name__ == "Shape" and not shape.text_frame.text.strip():
            try:
                fill = rgb(shape.fill.fore_color, (240, 240, 240))
            except Exception:
                fill = (240, 240, 240)
            box = [px(shape.left), px(shape.top),
                   px(shape.left) + px(shape.width),
                   px(shape.top) + px(shape.height)]
            radius = 12 if px(shape.height) > 30 else 3
            rounded(draw, box, radius=radius, fill=fill)
            continue
        overflow = draw_text_frame(draw, shape)
        if overflow and overflow > 6:
            warnings.append("text overflows its box by %d px: %r" %
                            (overflow, shape.text_frame.text[:48]))

    path = os.path.join(OUT, "slide_%02d.png" % index)
    canvas.save(path)
    return path, warnings


def main():
    os.makedirs(OUT, exist_ok=True)
    prs = Presentation(DECK)
    wanted = {int(a) for a in sys.argv[1:]} if len(sys.argv) > 1 else None
    problems = 0
    for index, slide in enumerate(prs.slides, start=1):
        if wanted and index not in wanted:
            continue
        path, warnings = render(slide, index)
        flag = "  <-- CHECK" if warnings else ""
        print("slide %02d  %-28s %s%s" %
              (index, os.path.basename(path),
               "%d shape(s)" % len(slide.shapes), flag))
        for warning in warnings:
            problems += 1
            print("          %s" % warning)
    print("\n%d potential layout problem(s)" % problems)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

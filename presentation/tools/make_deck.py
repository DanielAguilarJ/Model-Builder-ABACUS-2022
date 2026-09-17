# -*- coding: utf-8 -*-
"""Build the Model Builder 4.2 presentation (16:9, Keynote-style).

All figures come from /projects/sandbox/assets (real screenshots, real Abaqus
renders, diagrams rendered from live engine values).

The deck exists in Spanish and English from one source: L(es, en) picks the
wording and MB_LANG picks the language, so the layout is fixed in one place.

    MB_LANG=es python3 make_deck.py
    MB_LANG=en python3 make_deck.py
"""
from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ASSETS = "/projects/sandbox/assets"
LANG = os.environ.get("MB_LANG", "es")


def L(es_text, en_text):
    """Pick the wording for the active language."""
    return en_text if LANG == "en" else es_text


DIAGRAMS = os.path.join(ASSETS, "trimmed" if LANG == "es" else "trimmed_%s" % LANG)
FRAMED = os.path.join(ASSETS, "framed" if LANG == "es" else "framed_%s" % LANG)
OUTPUT = os.path.join(
    "/projects/sandbox/Model-Builder-ABACUS-2022/presentation",
    L("ModelBuilder_4.2_Presentacion.pptx", "ModelBuilder_4.2_Presentation.pptx"))

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_W = Inches(13.333 - 1.7)

FONT = "Arial"
INK = RGBColor(0x1D, 0x1D, 0x1F)
SOFT = RGBColor(0x6E, 0x6E, 0x73)
MUTED = RGBColor(0x86, 0x86, 0x8B)
LINE = RGBColor(0xD2, 0xD2, 0xD7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NEAR_WHITE = RGBColor(0xF5, 0xF5, 0xF7)
BLACK = RGBColor(0x0A, 0x0A, 0x0C)
BLUE = RGBColor(0x00, 0x71, 0xE3)
BLUE_DARK = RGBColor(0x29, 0x97, 0xFF)
GREEN = RGBColor(0x1B, 0x5E, 0x20)
AMBER = RGBColor(0x8A, 0x4B, 0x08)
RED = RGBColor(0xB7, 0x1C, 0x1C)
GREY_TEXT = RGBColor(0x9E, 0x9E, 0xA7)


# ---------------------------------------------------------------- primitives
def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_slide(prs, background=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background
    return slide


def textbox(slide, left, top, width, height, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.paragraphs[0].alignment = align
    return box, frame


def write(frame, blocks, first_paragraph=None):
    for index, block in enumerate(blocks):
        if index == 0 and first_paragraph is not None:
            paragraph = first_paragraph
        elif index == 0:
            paragraph = frame.paragraphs[0]
        else:
            paragraph = frame.add_paragraph()
        paragraph.alignment = block.get("align", PP_ALIGN.LEFT)
        if block.get("space_before") is not None:
            paragraph.space_before = Pt(block["space_before"])
        if block.get("space_after") is not None:
            paragraph.space_after = Pt(block["space_after"])
        paragraph.line_spacing = block.get("spacing", 1.0)
        run = paragraph.add_run()
        run.text = block["text"]
        font = run.font
        font.name = block.get("font", FONT)
        font.size = Pt(block["size"])
        font.bold = block.get("bold", False)
        font.color.rgb = block.get("color", INK)
    return frame


def rule(slide, left, top, width, color=LINE, weight=1.0):
    """A hairline separator. weight is in points."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width,
                                   Pt(weight))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def panel(slide, left, top, width, height, color=NEAR_WHITE, radius=0.035):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top,
                                   width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.adjustments[0] = radius
    return shape


def picture_fit(slide, path, box_left, box_top, box_width, box_height,
                align="center"):
    """Place an image scaled to fit the box, preserving the aspect ratio."""
    with Image.open(path) as image:
        aspect = image.size[0] / float(image.size[1])
    box_aspect = box_width / float(box_height)
    if aspect >= box_aspect:
        width = box_width
        height = Emu(int(box_width / aspect))
    else:
        height = box_height
        width = Emu(int(box_height * aspect))
    left = box_left if align == "left" else Emu(
        int(box_left + (box_width - width) / 2))
    top = Emu(int(box_top + (box_height - height) / 2))
    return slide.shapes.add_picture(path, left, top, width=width, height=height)


def slide_number(slide, index, dark=False):
    box, frame = textbox(slide, Emu(int(SLIDE_W - MARGIN - Inches(1.0))),
                         Inches(6.94), Inches(1.0), Inches(0.3),
                         align=PP_ALIGN.RIGHT)
    write(frame, [{"text": "%02d" % index, "size": 10,
                   "color": GREY_TEXT if dark else MUTED}])
    return box


def kicker(slide, text, dark=False, color=None):
    box, frame = textbox(slide, MARGIN, Inches(0.62), CONTENT_W, Inches(0.32))
    write(frame, [{"text": text.upper(), "size": 11.5, "bold": True,
                   "color": color or (BLUE_DARK if dark else BLUE)}])
    return box


def heading(slide, text, top=Inches(1.02), size=33, color=None, width=None):
    box, frame = textbox(slide, MARGIN, top, width or CONTENT_W, Inches(1.0))
    write(frame, [{"text": text, "size": size, "bold": True,
                   "color": color or INK, "spacing": 0.95}])
    return box


def subheading(slide, text, top=Inches(1.86), size=14.5, color=None,
               width=None):
    box, frame = textbox(slide, MARGIN, top, width or Inches(10.4), Inches(0.9))
    write(frame, [{"text": text, "size": size, "color": color or SOFT,
                   "spacing": 1.25}])
    return box


def caption(slide, text, top=Inches(6.62), size=11.5, color=None, width=None):
    box, frame = textbox(slide, MARGIN, top, width or Inches(11.0), Inches(0.7))
    write(frame, [{"text": text, "size": size, "color": color or MUTED,
                   "spacing": 1.25}])
    return box


# -------------------------------------------------------------- slide recipes
def slide_hero(prs, index=1):
    slide = add_slide(prs, BLACK)
    box, frame = textbox(slide, MARGIN, Inches(2.18), Inches(11.0), Inches(3.2))
    write(frame, [
        {"text": "Model Builder 4.2", "size": 62, "bold": True, "color": WHITE,
         "spacing": 0.92},
        {"text": L("Automatización trazable de uniones eje–cubo con chaveta\n"
                   "para Abaqus 2022",
                   "Traceable automation of shaft–hub key joints\n"
                   "for Abaqus 2022"),
         "size": 21, "color": GREY_TEXT, "spacing": 1.3, "space_before": 16},
    ])
    box2, frame2 = textbox(slide, MARGIN, Inches(6.28), Inches(11.5), Inches(0.8))
    write(frame2, [
        {"text": L("DIN 6885-1  ·  DIN 6892 Métodos A / B / C  ·  FVA 600 III",
                   "DIN 6885-1  ·  DIN 6892 Methods A / B / C  ·  FVA 600 III"),
         "size": 12.5, "color": RGBColor(0x7C, 0x7C, 0x86), "bold": True},
    ])
    slide_number(slide, index, dark=True)
    return slide


def slide_problema(prs, index=2):
    slide = add_slide(prs, BLACK)
    kicker(slide, L("el problema", "the problem"), dark=True)
    box, frame = textbox(slide, MARGIN, Inches(1.55), Inches(11.4), Inches(3.4))
    write(frame, [
        {"text": L("Un modelo de chavetero bien hecho\nno es difícil: es "
                   "difícil de repetir.",
                   "A good keyway model is not the hard part.\nRepeating it "
                   "is."),
         "size": 40, "bold": True, "color": WHITE, "spacing": 1.05},
    ])
    items = [
        (L("Cada modelo se rehacía a mano", "Every model was rebuilt by hand"),
         L("la tabla DIN se consultaba en papel y se teclaba caso por caso",
           "the DIN table was read on paper and typed in case by case")),
        (L("La malla dependía del autor", "The mesh depended on its author"),
         L("sin criterio de aceptación explícito ni registro de intentos",
           "with no explicit acceptance criterion and no record of attempts")),
        (L("El resultado no se podía auditar", "The result could not be audited"),
         L("sin saber qué versión, qué norma y qué hipótesis lo produjeron",
           "with no way to know which version, standard and assumptions "
           "produced it")),
    ]
    for index, (title, body) in enumerate(items):
        left = Emu(int(MARGIN + index * Inches(3.83)))
        box, frame = textbox(slide, left, Inches(5.15), Inches(3.55), Inches(1.6))
        write(frame, [
            {"text": title, "size": 15, "bold": True, "color": WHITE,
             "spacing": 1.15},
            {"text": body, "size": 12, "color": GREY_TEXT, "spacing": 1.3,
             "space_before": 7},
        ])
        rule(slide, left, Inches(4.86), Inches(0.9), color=BLUE_DARK, weight=2.0)
    slide_number(slide, index, dark=True)
    return slide


def slide_kpis(prs, index=3):
    slide = add_slide(prs)
    kicker(slide, L("qué es", "what it is"))
    heading(slide, L("Una sola herramienta desde la tabla de la norma\nhasta el "
                     "informe firmado",
                     "One tool, from the table in the standard\nto the signed "
                     "report"), size=31)
    kpis = [
        ("14", L("módulos de Python", "Python modules"),
         L("escritorio, kernel de Abaqus\ny núcleo compartido",
           "desktop, Abaqus kernel\nand the shared core")),
        ("26", L("filas DIN 6885-1", "DIN 6885-1 rows"),
         L("transcritas sin extrapolar,\n6 < d1 ≤ 500 mm",
           "transcribed with no extrapolation,\n6 < d1 ≤ 500 mm")),
        ("8", L("plantillas de malla", "mesh templates"),
         L("con puertas de calidad duras\ny presupuesto de elementos",
           "with hard quality gates\nand element budgets")),
        ("4", L("métodos DIN 6892", "DIN 6892 methods"),
         L("A, B, B-FVA y C, calculados\nen paralelo",
           "A, B, B-FVA and C, evaluated\nin parallel")),
        ("3", L("idiomas", "languages"),
         L("español, inglés y alemán\n650 claves con paridad",
           "Spanish, English and German\n650 keys at parity")),
    ]
    top = Inches(2.95)
    width = Inches(2.12)
    gap = Inches(0.28)
    for index, (number, label, note) in enumerate(kpis):
        left = Emu(int(MARGIN + index * (width + gap)))
        rule(slide, left, top, Inches(1.5), color=LINE)
        box, frame = textbox(slide, left, Emu(int(top + Inches(0.28))), width,
                             Inches(2.6))
        write(frame, [
            {"text": number, "size": 46, "bold": True, "color": INK,
             "spacing": 0.9},
            {"text": label, "size": 13.5, "bold": True, "color": INK,
             "spacing": 1.15, "space_before": 4},
            {"text": note, "size": 11, "color": MUTED, "spacing": 1.3,
             "space_before": 6},
        ])
    caption(slide, L("Todo lo que se ve en estas diapositivas sale de ejecutar "
                     "el código: las capturas son de la aplicación real y las "
                     "cifras las calcula el motor empaquetado.",
                     "Everything on these slides comes from running the code: "
                     "the screenshots are of the real application and the "
                     "figures are computed by the packaged engine."),
            top=Inches(6.45))
    slide_number(slide, index)
    return slide


def slide_diagram(prs, index, image, kicker_text=None, caption_text=None,
                  background=WHITE, top=Inches(1.28), bottom=Inches(6.45)):
    slide = add_slide(prs, background)
    if kicker_text:
        kicker(slide, kicker_text)
    box_top = top if kicker_text else Inches(0.95)
    picture_fit(slide, os.path.join(DIAGRAMS, image), MARGIN, box_top,
                CONTENT_W, Emu(int(bottom - box_top)))
    if caption_text:
        caption(slide, caption_text, top=Inches(6.58))
    slide_number(slide, index)
    return slide


def slide_shot(prs, index, image, kicker_text, title, body, caption_text=None):
    slide = add_slide(prs)
    kicker(slide, kicker_text)
    text_w = Inches(4.25)
    box, frame = textbox(slide, MARGIN, Inches(1.12), text_w, Inches(4.8))
    blocks = [{"text": title, "size": 26, "bold": True, "color": INK,
               "spacing": 1.0}]
    for paragraph in body:
        blocks.append({"text": paragraph, "size": 12.5, "color": SOFT,
                       "spacing": 1.35, "space_before": 12})
    write(frame, blocks)
    picture_fit(slide, os.path.join(FRAMED, image), Inches(5.32), Inches(0.72),
                Inches(7.45), Inches(6.06))
    if caption_text:
        box2, frame2 = textbox(slide, MARGIN, Inches(6.35), text_w, Inches(0.8))
        write(frame2, [{"text": caption_text, "size": 10.5, "color": MUTED,
                        "spacing": 1.3}])
    slide_number(slide, index)
    return slide


def slide_render(prs, index, image, kicker_text, title, body, dark=True):
    slide = add_slide(prs, BLACK if dark else WHITE)
    kicker(slide, kicker_text, dark=dark)
    box, frame = textbox(slide, MARGIN, Inches(1.12), Inches(4.2), Inches(4.6))
    blocks = [{"text": title, "size": 28, "bold": True,
               "color": WHITE if dark else INK, "spacing": 1.0}]
    for paragraph in body:
        blocks.append({"text": paragraph, "size": 12.5,
                       "color": GREY_TEXT if dark else SOFT, "spacing": 1.35,
                       "space_before": 13})
    write(frame, blocks)
    picture_fit(slide, os.path.join(FRAMED, image), Inches(5.35), Inches(0.75),
                Inches(7.4), Inches(6.0))
    slide_number(slide, index, dark=dark)
    return slide


def slide_din6885(prs, index):
    slide = add_slide(prs)
    kicker(slide, L("el dato normativo", "the normative datum"))
    heading(slide, L("La tabla no se interpola: o hay fila, o hay error",
                     "The table is never interpolated: either a row or an "
                     "error"), size=31)
    subheading(slide, L("din6885_row() devuelve la fila exacta con su etiqueta "
                        "NORMATIVE. Fuera del rango de la norma lanza una "
                        "excepción, para que nadie pueda declarar conformidad "
                        "DIN con la fila más parecida.",
                        "din6885_row() returns the exact row with its "
                        "NORMATIVE badge. Outside the range of the standard it "
                        "raises, so nobody can claim DIN compliance from the "
                        "nearest row."))
    cards = [
        ("26", L("filas de sección", "cross-section rows"),
         L("b × h, t1, t2, tolerancia de t1 y el incremento d2 para cada banda "
           "de diámetro",
           "b × h, t1, t2, the t1 tolerance and the d2 increment for every "
           "diameter band")),
        ("34", L("longitudes normalizadas", "standard lengths"),
         L("de 6 a 400 mm; la chaveta se ajusta hacia abajo, nunca hacia "
           "arriba",
           "from 6 to 400 mm; the key is snapped down, never up")),
        ("7", L("bandas de radios", "radius bands"),
         L("r1 de boca y r2 de fondo por separado: no son el mismo radio",
           "mouth r1 and root r2 separately: they are not the same radius")),
        ("3", L("formas implementadas", "forms implemented"),
         L("A, B y AB. De C a J se rechazan con un mensaje propio, no en "
           "silencio",
           "A, B and AB. C through J are rejected with their own message, not "
           "silently")),
    ]
    top = Inches(3.05)
    height = Inches(1.78)
    width = Inches(2.78)
    gap = Inches(0.17)
    for index_card, (number, label, note) in enumerate(cards):
        left = Emu(int(MARGIN + index_card * (width + gap)))
        panel(slide, left, top, width, height)
        box, frame = textbox(slide, Emu(int(left + Inches(0.26))),
                             Emu(int(top + Inches(0.22))),
                             Emu(int(width - Inches(0.5))), Inches(1.42))
        write(frame, [
            {"text": number + "   " + label, "size": 14.5, "bold": True,
             "color": INK, "spacing": 1.05},
            {"text": note, "size": 10.5, "color": SOFT, "spacing": 1.3,
             "space_before": 6},
        ])
    box, frame = textbox(slide, MARGIN, Inches(5.18), CONTENT_W, Inches(1.4))
    write(frame, [
        {"text": L("D40, la fila de referencia del proyecto",
                   "D40, the project's reference row"),
         "size": 15, "bold": True, "color": INK},
        {"text": L("38 < d1 ≤ 44 mm     b × h = 12 × 8 mm     "
                   "t1 = 5,00 +0,20 mm     t2 = 3,30 mm     "
                   "d2(ref) = 48,00 mm     r1 = 0,25…0,40 mm     "
                   "r2 = 0,16…0,25 mm",
                   "38 < d1 ≤ 44 mm     b × h = 12 × 8 mm     "
                   "t1 = 5.00 +0.20 mm     t2 = 3.30 mm     "
                   "d2(ref) = 48.00 mm     r1 = 0.25…0.40 mm     "
                   "r2 = 0.16…0.25 mm"),
         "size": 13, "color": SOFT, "font": "Consolas", "space_before": 10,
         "spacing": 1.35},
    ])
    caption(slide, L("El chaflán de la chaveta no está en la tabla: es un "
                     "parámetro del usuario (c = 0,8 mm por defecto) y así se "
                     "declara en el informe.",
                     "The key chamfer is not in the table: it is a user "
                     "parameter (c = 0.8 mm by default) and the report says "
                     "so."), top=Inches(6.5))
    slide_number(slide, index)
    return slide


def slide_verificacion(prs, index):
    slide = add_slide(prs)
    kicker(slide, L("verificación", "verification"))
    heading(slide, L("Lo que se ha comprobado, y cómo",
                     "What has been checked, and how"), size=33)
    left_items = [
        (L("Autotest empaquetado", "Packaged self-test"),
         L("identidad 4.2 y esquema 6, migración 5→6, catálogos DIN y FVA, "
           "cálculos de referencia, ocho plantillas de malla, informes en los "
           "tres idiomas. Termina con código 0.",
           "4.2 identity and schema 6, the 5→6 migration, DIN and FVA "
           "catalogues, golden calculations, eight mesh templates, reports in "
           "all three languages. Exits with code 0.")),
        (L("Paridad de los tres catálogos", "Parity across the three catalogues"),
         L("650 claves por idioma, sin huecos; cada código de validación tiene "
           "su presentación traducida.",
           "650 keys per language with no gaps; every validation code has its "
           "translated presentation.")),
        (L("Configuraciones vigiladas", "Guarded configurations"),
         L("los valores por defecto y los dos compañeros FVA empaquetados se "
           "validan bajo esquema 6. La configuración de 10 ciclos del "
           "repositorio se valida por separado y no se empaqueta.",
           "the defaults and the two packaged FVA companions validate under "
           "schema 6. The repository's 10-cycle setup is validated separately "
           "and is not packaged.")),
    ]
    right_items = [
        (L("Compatibilidad con el kernel", "Compatibility with the kernel"),
         L("los seis módulos que corren dentro de Abaqus se auditaron con AST: "
           "sin f-strings, sin anotaciones, sólo os, sys, json, math, time y "
           "hashlib. División importada del futuro.",
           "the six modules that run inside Abaqus were audited with AST: no "
           "f-strings, no annotations, only os, sys, json, math, time and "
           "hashlib. Division imported from __future__.")),
        (L("Contrato de artefactos", "Artifact contract"),
         L("el proyecto se declara correcto sólo si existen la auditoría, "
           "BUILD_RESULT y todos los artefactos solicitados o requeridos, y si "
           "la malla realizada coincide con la planificada.",
           "a project is declared sound only if the audit, BUILD_RESULT and all "
           "requested or required artifacts exist, and the realised mesh "
           "matches the planned one.")),
        (L("Cifras fijadas en el test", "Figures locked in the test"),
         L("999,18 N·m gobernado por el cubo, t2tr = 3,141 mm, f_W de la "
           "ecuación 3 y el dominio de K_λ.",
           "999.18 N·m hub-governed, t2tr = 3.141 mm, f_W from equation 3 and "
           "the K_λ domain.")),
    ]
    for column, items in enumerate((left_items, right_items)):
        left = Emu(int(MARGIN + column * Inches(5.95)))
        top = Inches(2.35)
        for title, body in items:
            rule(slide, left, top, Inches(0.75), color=GREEN, weight=2.2)
            box, frame = textbox(slide, left, Emu(int(top + Inches(0.2))),
                                 Inches(5.4), Inches(1.4))
            write(frame, [
                {"text": title, "size": 15, "bold": True, "color": INK,
                 "spacing": 1.1},
                {"text": body, "size": 11.5, "color": SOFT, "spacing": 1.32,
                 "space_before": 6},
            ])
            top = Emu(int(top + Inches(1.42)))
    caption(slide, L("Lo que el autotest no hace, y así lo declara: no arranca "
                     "Abaqus, no resuelve, no compila LaTeX y no valida un "
                     "especimen FVA.",
                     "What the self-test does not do, and says so: it does not "
                     "start Abaqus, does not solve, does not compile LaTeX and "
                     "does not validate an FVA specimen."), top=Inches(6.72))
    slide_number(slide, index)
    return slide


def slide_pasos(prs, index):
    slide = add_slide(prs, BLACK)
    kicker(slide, L("próximos pasos", "next steps"), dark=True)
    heading(slide, L("Tres decisiones y una compilación",
                     "Three decisions and one build"), size=36, color=WHITE)
    steps = [
        ("01", L("Reconstruir el ejecutable", "Rebuild the executable"),
         L("make_exe.bat, sin cambios, desde el árbol completo del proyecto. "
           "El .exe publicado es anterior a las dos correcciones: 17 de 31 "
           "entradas del manifiesto han cambiado.",
           "make_exe.bat, unchanged, from the full project tree. The published "
           ".exe predates both corrections: 17 of the 31 manifest entries have "
           "changed.")),
        ("02", L("Fijar el signo de la ecuación 9",
                 "Settle the sign of equation 9"),
         L("contra el texto DIN 6892 con licencia. El signo cambia un 46,35 % "
           "la profundidad y el componente del eje; en el D40 por defecto el "
           "cubo sigue gobernando a 999,18 N·m.",
           "against the licensed DIN 6892 text. The sign changes the shaft "
           "depth and shaft component by 46.35%; in the bundled D40 default, "
           "the hub still governs at 999.18 N·m.")),
        ("03", L("Numerar la versión", "Number the version"),
         L("dos ejecutables que se identifican como 4.2 y devuelven pares "
           "distintos son una colisión de identidad. Mi recomendación es 4.3.",
           "two executables that both report 4.2 and return different torques "
           "are an identity collision. My recommendation is 4.3.")),
    ]
    top = Inches(2.7)
    for number, title, body in steps:
        box, frame = textbox(slide, MARGIN, top, Inches(1.0), Inches(0.6))
        write(frame, [{"text": number, "size": 17, "bold": True,
                       "color": BLUE_DARK}])
        box2, frame2 = textbox(slide, Inches(1.95), top, Inches(10.5), Inches(1.3))
        write(frame2, [
            {"text": title, "size": 19, "bold": True, "color": WHITE,
             "spacing": 1.05},
            {"text": body, "size": 12.5, "color": GREY_TEXT, "spacing": 1.35,
             "space_before": 7},
        ])
        rule(slide, MARGIN, Emu(int(top - Inches(0.24))), CONTENT_W,
             color=RGBColor(0x2A, 0x2A, 0x30))
        top = Emu(int(top + Inches(1.42)))
    slide_number(slide, index, dark=True)
    return slide


def slide_cierre(prs, index):
    slide = add_slide(prs, BLACK)
    box, frame = textbox(slide, MARGIN, Inches(2.5), Inches(11.0), Inches(2.6))
    write(frame, [
        {"text": L("El valor no está en el modelo.\nEstá en poder demostrar "
                   "cómo se hizo.",
                   "The value is not the model.\nIt is being able to show how "
                   "it was made."),
         "size": 36, "bold": True, "color": WHITE, "spacing": 1.1},
    ])
    box2, frame2 = textbox(slide, MARGIN, Inches(5.6), Inches(11.0), Inches(1.2))
    write(frame2, [
        {"text": L("Model Builder 4.2  ·  esquema de parámetros 6  ·  "
                   "din6892-methods-1.0  ·  auto-mesh-2.0",
                   "Model Builder 4.2  ·  parameter schema 6  ·  "
                   "din6892-methods-1.0  ·  auto-mesh-2.0"),
         "size": 12.5, "color": GREY_TEXT},
        {"text": L("Cada cálculo se puede rastrear al código; las capturas y "
                   "renders históricos se identifican como evidencia de origen.",
                   "Every calculation can be traced to code; historical "
                   "screenshots and renders are identified as source evidence."),
         "size": 12.5, "color": RGBColor(0x7C, 0x7C, 0x86), "space_before": 8},
    ])
    slide_number(slide, index, dark=True)
    return slide



def slide_panels(prs, index, kicker_text, heading_text, sub_text, cards,
                 caption_text=None, columns=3, heading_size=30):
    """A grid of small titled panels: used for the feature detail slides."""
    slide = add_slide(prs)
    kicker(slide, kicker_text)
    heading(slide, heading_text, size=heading_size)
    if sub_text:
        subheading(slide, sub_text)
    rows = (len(cards) + columns - 1) // columns
    gap = Inches(0.22)
    width = Emu(int((CONTENT_W - gap * (columns - 1)) / columns))
    top0 = Inches(2.72 if sub_text else 2.20)
    height = Inches(1.82 if rows > 1 else 2.10)
    row_gap = Inches(0.20)
    for position, (title, body) in enumerate(cards):
        column = position % columns
        row = position // columns
        left = Emu(int(MARGIN + column * (width + gap)))
        top = Emu(int(top0 + row * (height + row_gap)))
        panel(slide, left, top, width, height)
        box, frame = textbox(slide, Emu(int(left + Inches(0.24))),
                             Emu(int(top + Inches(0.20))),
                             Emu(int(width - Inches(0.48))),
                             Emu(int(height - Inches(0.32))))
        write(frame, [
            {"text": title, "size": 13.5, "bold": True, "color": INK,
             "spacing": 1.1},
            {"text": body, "size": 10.5, "color": SOFT, "spacing": 1.3,
             "space_before": 6},
        ])
    if caption_text:
        caption(slide, caption_text, top=Inches(6.6))
    slide_number(slide, index)
    return slide


# ------------------------------------------------------------------ assembly
def build():
    prs = new_deck()
    counter = [0]

    def N():
        counter[0] += 1
        return counter[0]

    slide_hero(prs, N())
    slide_problema(prs, N())
    slide_kpis(prs, N())

    slide_diagram(prs, N(), "dia_arquitectura.png",
                  L("arquitectura", "architecture"),
                  L("Un cambio en la tabla DIN o en un factor de DIN 6892 se "
                    "escribe una sola vez y llega a los dos intérpretes.",
                    "A change to the DIN table or to a DIN 6892 factor is "
                    "written once and reaches both interpreters."))
    slide_diagram(prs, N(), "dia_tabs.png", L("cómo se usa", "how it is used"),
                  L("Seis pestañas en el orden en que se toman las decisiones: "
                    "primero el proyecto, después la geometría normativa, y "
                    "sólo al final la ejecución.",
                    "Six tabs in the order the decisions are actually made: "
                    "the project first, then the normative geometry, and only "
                    "at the end the run."),
                  top=Inches(2.3), bottom=Inches(5.4))

    shots = [
        (6, "gui_01_proyecto.png", L("pestaña 1", "tab 1"),
         L("Proyecto", "Project"),
         [L("Nombre, autor e idioma quedan grabados en project.json: el "
            "informe se emite en el idioma del proyecto, no en el de quien lo "
            "abre.",
            "Name, author and language are recorded in project.json: the "
            "report is issued in the project's language, not in the language "
            "of whoever opens it."),
          L("La ruta de Abaqus se detecta buscando los lanzadores instalados, "
            "sin ejecutar ninguno.",
            "The Abaqus path is detected by looking for installed launchers, "
            "without executing any of them."),
          L("La carpeta raíz por defecto es Documentos/ModelBuilder Projects.",
            "The default root folder is Documents/ModelBuilder Projects.")],
         L("Captura real de la aplicación, idioma por defecto español.",
           "Real capture of the application, with the interface switched to "
           "English.")),
        (7, "gui_02_geometria_din.png", L("pestaña 2", "tab 2"),
         L("Geometría DIN", "DIN geometry"),
         [L("A la izquierda, los pocos parámetros independientes que el "
            "usuario decide. A la derecha, todo lo que la norma deriva de "
            "ellos.",
            "On the left, the few independent parameters the user decides. On "
            "the right, everything the standard derives from them."),
          L("El panel derecho se recalcula al teclear: fila DIN, longitudes, "
            "cubo, plan de malla y la lista de validación completa.",
            "The right panel recomputes as you type: DIN row, lengths, hub, "
            "mesh plan and the complete validation list."),
          L("La etiqueta NORMATIVE en verde indica que la geometría está "
            "dentro del sobre de la norma.",
            "The green NORMATIVE badge means the geometry sits inside the "
            "envelope of the standard.")],
         L("d1 = 40 mm  ·  banda 38 < d1 ≤ 44  ·  b × h = 12 × 8  ·  "
           "DIN 6885-1 Form A 12 × 8 × 50.",
           "d1 = 40 mm  ·  band 38 < d1 ≤ 44  ·  b × h = 12 × 8  ·  "
           "DIN 6885-1 Form A 12 × 8 × 50.")),
        (8, "gui_03_malla_materiales.png", L("pestaña 3", "tab 3"),
         L("Malla y materiales", "Mesh and materials"),
         [L("La plantilla se elige por nombre; el modo AUTO compara candidatos "
            "completos por pieza y conserva el ganador.",
            "The template is chosen by name; AUTO mode compares complete "
            "candidates per part and keeps the winner."),
          L("Una semilla en 0 significa automático, y la regla sigue al "
            "refinamiento: con la banda de entalla activa son D/16 en el "
            "eje y d_a/40 en el cubo. La chaveta siempre b/20.",
            "A seed left at 0 means automatic, and the rule follows the "
            "refinement: with the notch band on it is D/16 on the shaft "
            "and d_a/40 on the hub. The key is always b/20."),
          L("El criterio rápido de presión está marcado CORE-SCREENING: es "
            "plausibilidad, no DIN 6892 Método B.",
            "The quick pressure check is badged CORE-SCREENING: it is "
            "plausibility, not DIN 6892 Method B.")],
         L("Aquí la plantilla es HEX_CERTIFIED, en modo auto y con 3 intentos.",
           "Here the template is HEX_CERTIFIED, in auto mode with 3 "
           "attempts.")),
        (9, "gui_04_analisis.png", L("pestaña 4", "tab 4"),
         L("Análisis", "Analysis"),
         [L("El análisis es opcional y está desactivado por defecto: construir "
            "y auditar no requiere resolver.",
            "The analysis is optional and off by default: building and "
            "auditing does not require solving."),
          L("Si se activa, se crea contacto general, un step estático TORSION "
            "y un momento sobre el punto de referencia del extremo motriz.",
            "When enabled it creates general contact, a static TORSION step "
            "and a moment on the reference point of the drive end."),
          L("El estudio de convergencia de malla se lanza aparte y escribe su "
            "propio CSV.",
            "The mesh convergence study runs separately and writes its own "
            "CSV.")],
         L("Crear el job y enviarlo al solver son dos casillas distintas.",
           "Creating the job and submitting it to the solver are two separate "
           "checkboxes.")),
        (10, "gui_05_din6892_fva.png", L("pestaña 5", "tab 5"),
         "DIN 6892 / FVA",
         [L("El método se elige explícitamente. Los factores licenciados que "
            "no se pueden transcribir se quedan en 1,0 y el resultado se marca "
            "PROVISIONAL_FACTORS.",
            "The method is chosen explicitly. Licensed factors that cannot be "
            "transcribed stay at 1.0 and the result is flagged "
            "PROVISIONAL_FACTORS."),
          L("N_W en blanco significa deducir de la razón de carga R; con un "
            "par pulsante da cero inversiones y f_W = 1.",
            "N_W left blank means derive it from the load ratio R; a "
            "pulsating torque gives zero reversals and f_W = 1."),
          L("El Método A no se calcula aquí: se importa un CSV de un ODB ya "
            "resuelto, con SHA-256 y cobertura verificadas.",
            "Method A is not computed here: a CSV from an already solved ODB "
            "is imported, with SHA-256 and coverage verified.")],
         L("Las normas seleccionadas se declaran una por una, incluidas las "
           "que no están implementadas.",
           "The selected standards are declared one by one, including the ones "
           "that are not implemented.")),
        (11, "gui_06_ejecutar_evidencias.png", L("pestaña 6", "tab 6"),
         L("Ejecutar y evidencias", "Run and evidence"),
         [L("El log de Abaqus se transmite en vivo a la ventana y al archivo "
            "logs/abaqus.log, con el comando y el directorio de trabajo en la "
            "cabecera.",
            "The Abaqus log is streamed live to the window and to "
            "logs/abaqus.log, with the command and the working directory in "
            "the header."),
          L("Al terminar, la lista de artefactos se rellena con tamaño y "
            "etiqueta de evidencia; los botones abren el proyecto, el informe "
            "y el modelo.",
            "When it finishes, the artifact list fills in with size and "
            "evidence badge; the buttons open the project, the report and the "
            "model."),
          L("Detener mata el árbol de procesos del solver y marca el proyecto "
            "como CANCELLED.",
            "Stop kills the solver process tree and marks the project "
            "CANCELLED.")],
         L("Ningún archivo de solver cae junto al ejecutable: el subproceso "
           "corre con cwd = jobs/.",
           "No solver file lands next to the executable: the subprocess runs "
           "with cwd = jobs/.")),
        (12, "gui_07_validacion.png",
         L("control de calidad", "quality control"),
         L("La validación habla en códigos", "Validation speaks in codes"),
         [L("Cada aviso tiene un código estable, un mensaje canónico en inglés "
            "y una presentación traducida. El código es el que se cita en el "
            "informe y en una revisión.",
            "Every issue has a stable code, a canonical English message and a "
            "translated presentation. The code is what gets cited in the "
            "report and in a review."),
          L("El diálogo informa el resultado, no lo adorna: sin errores, el "
            "recuento de advertencias y el recordatorio de que el cálculo "
            "resistente sigue etiquetado CORE-SCREENING.",
            "The dialog reports the outcome, it does not dress it up: no "
            "errors, the warning count, and the reminder that the strength "
            "calculation stays labelled CORE-SCREENING."),
          L("Un error bloquea la construcción antes de tocar Abaqus; una "
            "advertencia informa y queda registrada en la auditoría.",
            "An error blocks the build before Abaqus is touched; a warning "
            "informs and is recorded in the audit.")],
         L("Por defecto: 0 errores y 3 advertencias — un factor licenciado "
           "todavía neutro, la ecuación 9 optimista y una relación de aspecto "
           "axial alta esperada en la entalla.",
           "By default: 0 errors and 3 warnings — a licensed factor still "
           "neutral, the optimistic equation 9, and a high axial aspect ratio "
           "expected at the notch.")),
    ]
    for args in shots:
        slide_shot(prs, N(), args[1], args[2], args[3], args[4], args[5])

    slide_diagram(prs, N(), "dia_pipeline.png",
                  L("flujo de ejecución", "execution flow"))
    slide_diagram(prs, N(), "dia_workspace.png",
                  L("lo que queda en disco", "what stays on disk"))

    slide_panels(
        prs, N(), L('puntos de partida', 'starting points'),
        L('Quince presets, ninguno que salte la validación',
          'Fifteen presets, none of which skips validation'),
        L('Un preset sólo fija parámetros. Después se fusiona sobre los valores por defecto y se vuelve a derivar y validar como si se hubiera teclado a mano.',
          'A preset only sets parameters. It is then merged onto the defaults and re-derived and re-validated exactly as if it had been typed in by hand.'),
        [
         (L('5 de geometría y análisis', '5 geometry and analysis'),
          L('D25 pequeña, D40 validada por defecto, D50 siguiente banda DIN, D40 cuadrática C3D20 y D40 con análisis de torsión.',
            'D25 small joint, D40 validated default, D50 next DIN band, D40 quadratic C3D20 and D40 with torsion analysis.')),
         (L('8 de malla', '8 mesh policies'),
          L('Una por plantilla: hex certificada, híbrida FVA Método A, equilibrada, calidad crítica, vista rápida, hex dominante, precisión cuadrática y respaldo robusto.',
            'One per template: hex certified, FVA Method A hybrid, balanced, quality critical, fast preview, hex dominant, quadratic accuracy and robust fallback.')),
         (L('2 de investigación FVA', '2 FVA research'),
          L('D40 pre-solve de 1 ciclo y configuración de Método A D40 de 20 ciclos con coincidencia a verificar durante el build. Las dos con el solver deliberadamente apagado.',
            'D40 1-cycle pre-solve and a 20-cycle D40 Method A setup whose matching must be verified at build time. Both with the solver deliberately off.')),
         (L('Un preset no es un atajo', 'A preset is not a shortcut'),
          L('Pasa por las mismas puertas: si el resultado sale del sobre normativo, el error bloquea igual.',
            'It goes through the same gates: if the result leaves the normative envelope, the error blocks it just the same.')),
         (L('Se guarda y se recupera', 'Saved and reloaded'),
          L('La configuración va y vuelve a JSON, así que un caso se entrega como archivo y se reproduce tal cual.',
            'The configuration round-trips to JSON, so a case can be handed over as a file and reproduced exactly.')),
         (L('Los presets del paquete están vigilados', 'The bundled presets are guarded'),
          L('El publicador comprueba campo por campo los dos compañeros FVA: no pueden cambiar en silencio entre versiones.',
            'The publisher checks the two FVA companions field by field: they cannot change silently between versions.')),
        ],
        L('Los presets de malla existen para hacer explícita la política, no para esconderla.',
          'The mesh presets exist to make the policy explicit, not to hide it.'),
        columns=3)

    slide_panels(
        prs, N(), L('espacio de diseño', 'design space'),
        L('Qué se puede construir, y qué se rechaza a propósito',
          'What can be built, and what is deliberately refused'),
        None,
        [
         (L('Formas de chaveta', 'Key forms'),
          L('A, B y AB implementadas. De C a J se rechazan con un mensaje propio: necesitan cotas de taladro y chaflán que no están implementadas.',
            'A, B and AB implemented. C through J are refused with their own message: they need hole and chamfer dimensions that are not implemented.')),
         (L('Cubo cilíndrico o cónico', 'Cylindrical or tapered hub'),
          L('El cónico añade un cuarto sólido, el casquillo, y desactiva el refinamiento de banda de la entalla; el modelo lo declara.',
            'Tapered adds a fourth solid, the bushing, and switches off the notch band refinement; the model declares it.')),
         (L('Diámetro', 'Diameter'),
          L('Continuo en 6 < d1 <= 500 mm, resuelto a una de las 26 filas. Fuera de ahí lanza excepción en lugar de aproximar.',
            'Continuous over 6 < d1 <= 500 mm, resolved to one of the 26 rows. Outside that it raises instead of approximating.')),
         (L('Orden y tipo de elemento', 'Element order and type'),
          L('Lineal o cuadrático (por defecto), C3D8/C3D8I o C3D20/C3D20R. C3D8R sólo se acepta como elección explícita heredada, y levanta advertencia.',
            'Linear or quadratic (default), C3D8/C3D8I or C3D20/C3D20R. C3D8R is accepted only as an explicit legacy choice, and it raises a warning.')),
         (L('Semillas automáticas', 'Automatic seeds'),
          L('Con refinamiento de entalla, que es el defecto: D/16 en el eje y d_a/40 en el cubo. Sin él, D/32 y d_a/80. La regla aplicada se escribe en la auditoría.',
            'With notch refinement, which is the default: D/16 on the shaft and d_a/40 on the hub. Without it, D/32 and d_a/80. The rule actually used is written into the audit.')),
         (L('Modo de cumplimiento', 'Compliance mode'),
          L('Normativo o anulación del usuario. Salirse del sobre está permitido; lo que no está permitido es que no se note.',
            'Normative or user override. Leaving the envelope is allowed; what is not allowed is for it to go unnoticed.')),
        ],
        None,
        columns=3)


    slide_render(prs, N(), "model_assembly_hex.png",
                 L("el resultado", "the result"),
                 L("Sólidos, no cascarones", "Solids, not shells"),
                 [L("Eje, chaveta y cubo se construyen como volúmenes "
                    "independientes y se ensamblan con las tolerancias de la "
                    "norma.",
                    "Shaft, key and hub are built as independent solids and "
                    "assembled with the tolerances of the standard."),
                  L("Esta evidencia muestra una malla hexaédrica. La "
                    "coincidencia exacta de interfaces no se deduce de la "
                    "imagen: el backend especializado de Método A la comprueba "
                    "comparando coordenadas de nodos después de mallar.",
                    "This evidence shows a hexahedral mesh. Exact interface "
                    "matching is not inferred from the image: the specialized "
                    "Method A backend checks it by comparing node coordinates "
                    "after meshing."),
                  L("La captura es la imagen que el propio motor guarda en "
                    "screenshots/ al terminar.",
                    "This is the image the engine itself saves to "
                    "screenshots/ when it finishes.")])
    slide_render(prs, N(), "model_notch_section.png",
                 L("el detalle que importa", "the detail that matters"),
                 L("La banda del chavetero", "The keyway band"),
                 [L("En la sección se ve la banda de elementos más finos que "
                    "rodea la ranura: no es una malla uniforme más densa, es "
                    "refinamiento donde está el gradiente.",
                    "The section shows the band of finer elements around the "
                    "keyway: not a uniformly denser mesh, but refinement where "
                    "the gradient is."),
                  L("El ancho de la banda y el número de elementos sobre el "
                    "arco del radio r2 son parámetros, y la auditoría "
                    "comprueba que se hayan cumplido.",
                    "The band width and the number of elements across the r2 "
                    "fillet arc are parameters, and the audit checks that they "
                    "were honoured."),
                  L("El conjunto de elementos NOTCH_SHAFT se construye "
                    "alrededor de las dos esquinas del chavetero para que la "
                    "consulta de concentración no se vaya a otro punto "
                    "caliente.",
                    "The NOTCH_SHAFT element set is built around both keyway "
                    "corners so the stress-concentration query cannot wander "
                    "off to an unrelated hot spot.")],
                 dark=False)

    slide_diagram(prs, N(), "dia_malla.png", L("política de malla", "mesh policy"))

    slide_diagram(prs, N(), "dia_mesh_loop.png",
                  L("selección de malla", "mesh selection"))

    slide_panels(
        prs, N(), L('montaje del análisis', 'analysis setup'),
        L('Si se pide análisis, esto es lo que se monta',
          'If the analysis is requested, this is what gets built'),
        None,
        [
         (L('Contacto', 'Contact'),
          L('Contacto general con rozamiento por penalización isótropo y comportamiento normal duro con separación permitida.',
            'General contact with isotropic penalty friction and hard normal behaviour with separation allowed.')),
         (L('Extremo motriz', 'Drive end'),
          L('Superficie en el extremo del eje, punto de referencia y acoplamiento cinemático sobre toda la superficie.',
            'Surface at the shaft end, a reference point and a kinematic coupling over the whole surface.')),
         (L('Región sujeta', 'Held region'),
          L('El diámetro exterior del cubo o sus dos caras, empotrado. Es una elección del usuario, no un supuesto oculto.',
            'The hub outer diameter or both hub faces, encastred. A user choice, not a hidden assumption.')),
         (L('Paso', 'Step'),
          L('Paso estático TORSION, con no linealidad geométrica y estabilización opcionales y control de incremento del formulario.',
            'Static TORSION step, with optional geometric nonlinearity and stabilization, and increment control from the form.')),
         (L('Carga', 'Load'),
          L('Momento sobre el eje aplicado en el punto de referencia; el punto guía bloquea todos los grados de libertad menos el giro.',
            'A moment about the shaft axis applied at the reference point; the guide point locks every degree of freedom except the rotation.')),
         (L('Salidas', 'Output'),
          L('Campo S, U, E, RF, CSTRESS y CDISP en cuatro intervalos; historia UR3 y RM3 en el punto de referencia.',
            'Field S, U, E, RF, CSTRESS and CDISP at four intervals; history UR3 and RM3 at the reference point.')),
        ],
        L('Crear el trabajo y enviarlo al solver siguen siendo dos decisiones separadas, y ninguna es el valor por defecto.',
          'Creating the job and submitting it to the solver remain two separate decisions, and neither is the default.'),
        columns=3)

    slide_diagram(prs, N(), "dia_badges.png", L("trazabilidad", "traceability"))
    slide_din6885(prs, N())

    slide_panels(
        prs, N(), L('la auditoría', 'the audit'),
        L('El archivo que dice qué se construyó de verdad',
          'The file that says what was actually built'),
        None,
        [
         (L('Dos formatos', 'Two formats'),
          L('PARAM_BUILD_AUDIT.txt para leerlo con los ojos y .json para procesarlo con una herramienta.',
            'PARAM_BUILD_AUDIT.txt to read with your eyes and .json to process with a tool.')),
         (L('Un veredicto', 'A verdict'),
          L('OK o CHECK, y el mismo veredicto se refleja en el resultado del build: no hay dos versiones de la verdad.',
            'OK or CHECK, and the same verdict is mirrored into the build result: there are not two versions of the truth.')),
         (L('Procedencia', 'Provenance'),
          L('SHA-256 de los archivos del motor y del núcleo que produjeron ese modelo, no de los que estén hoy en disco.',
            'SHA-256 of the engine and core files that produced that model, not of whatever is on disk today.')),
         (L('Plan frente a realidad', 'Plan versus reality'),
          L('La malla planificada y la realizada, pieza por pieza. Si no coinciden, el proyecto no se declara correcto.',
            'The planned mesh and the realised one, part by part. If they disagree, the project is not declared sound.')),
         (L('Por pieza', 'Per part'),
          L('Elementos, score de calidad, proporción de hexaedros, relación de aspecto, reparaciones e historial de intentos.',
            'Elements, quality score, hex share, aspect ratio, repairs and the attempt history.')),
         (L('La lista de avisos', 'The issue list'),
          L('Cada advertencia con su código estable, más las etiquetas de evidencia y los conjuntos y superficies creados.',
            'Every warning with its stable code, plus the evidence badges and the sets and surfaces created.')),
        ],
        None,
        columns=3)

    slide_panels(
        prs, N(), L('el entregable', 'the deliverable'),
        L('El informe es el producto, no un extra',
          'The report is the product, not an extra'),
        None,
        [
         (L('Diez secciones', 'Ten sections'),
          L('Proyecto y procedencia, alcance, geometría, DIN 6892 y normas, validación, malla y política, comprobaciones, evidencia, entregables y limitaciones.',
            'Project and provenance, scope, geometry, DIN 6892 and standards, validation, mesh and policy, checks, evidence, deliverables and limitations.')),
         (L('Tres archivos y un log', 'Three files and a log'),
          L('report.tex para revisar, report.pdf para firmar, report_data.json para procesar y latex_build.log para depurar.',
            'report.tex to review, report.pdf to sign, report_data.json to process and latex_build.log to debug.')),
         (L('Trilingüe', 'Trilingual'),
          L('Se emite en el idioma del proyecto, desde los mismos catálogos de 650 claves que usa la interfaz.',
            "Emitted in the project's language, from the same 650-key catalogues the interface uses.")),
         (L('Las etiquetas sobreviven', 'The badges survive'),
          L('Cada valor llega al PDF con su etiqueta de procedencia como distintivo de color, no como nota al pie.',
            'Every value reaches the PDF with its provenance badge as a coloured chip, not as a footnote.')),
         (L('Compilación', 'Compilation'),
          L('latexmk si está, dos pasadas de pdflatex si no, sin shell y con tiempo límite finito.',
            'latexmk when available, two pdflatex passes otherwise, with no shell and a finite timeout.')),
         (L('Sin LaTeX instalado', 'With no LaTeX installed'),
          L('Escribe el .tex, lo dice con un mensaje traducido y no tumba el build por algo que no es del modelo.',
            "It writes the .tex, says so with a translated message, and does not fail the build over something that is not the model's fault.")),
        ],
        None,
        columns=3)

    slide_diagram(prs, N(), "dia_metodos.png",
                  L("capacidad resistente", "load capacity"))
    slide_diagram(prs, N(), "chart_pares.png", L("resultados", "results"))
    slide_diagram(prs, N(), "dia_t2tr.png",
                  L("corrección 1 de 2", "correction 1 of 2"))
    slide_diagram(prs, N(), "chart_fw.png",
                  L("corrección 2 de 2", "correction 2 of 2"))
    slide_diagram(prs, N(), "dia_eq9.png", L("hallazgo abierto", "open finding"))

    slide_diagram(prs, N(), "dia_method_a.png", L("Método A", "Method A"))
    slide_diagram(prs, N(), "dia_matlab.png",
                  L("contraste numérico", "numerical cross-check"))

    slide_panels(
        prs, N(), L('catálogo FVA 600 III', 'FVA 600 III catalogue'),
        L('Ocho configuraciones concretas, no un dominio continuo',
          'Eight discrete configurations, not a continuous domain'),
        None,
        [
         (L('VB1 a VB8', 'VB1 to VB8'),
          L('Cada variante fija d_w, l_tr/d_w, Q_A, la interferencia, la razón de carga R y la forma de chaveta.',
            'Each variant fixes d_w, l_tr/d_w, Q_A, the interference, the load ratio R and the key form.')),
         (L('Son puntos, no un rango', 'They are points, not a range'),
          L('El catálogo está documentado como configuraciones discretas de investigación: interpolar entre ellas no está respaldado.',
            'The catalogue is documented as discrete research configurations: interpolating between them is not supported.')),
         (L('VB1, la referencia', 'VB1, the reference'),
          L('d_w = 40 mm, l_tr/d_w = 0,95, Q_A = 0,5, R = 0 y chaveta forma A. Es la variante de los presets del paquete.',
            'd_w = 40 mm, l_tr/d_w = 0.95, Q_A = 0.5, R = 0 and a form A key. It is the variant of the bundled presets.')),
         (L('Modelos de material', 'Material models'),
          L('Chaboche-Lemaitre combinado en el eje, UML Ramberg-Osgood en el cubo y elástico-plástico ideal en la chaveta.',
            'Chaboche-Lemaitre combined on the shaft, UML Ramberg-Osgood on the hub and elastic-ideal-plastic on the key.')),
         (L('Malla coincidente', 'Matching mesh'),
          L('Cinco interfaces verificadas después de mallar, con tamaño objetivo 0,8 mm y tolerancia de 1e-6 mm.',
            'Five interfaces verified after meshing, with a 0.8 mm target size and a 1e-6 mm tolerance.')),
         (L('COINCIDENTE se gana', 'MATCHED is earned'),
          L('Sólo se declara tras comparar coordenadas reales de nodos en cada interfaz que transmite carga. Sembrar igual no es prueba.',
            'Only declared after comparing real node coordinates on every load-carrying interface. Equal seeding is not proof.')),
        ],
        None,
        columns=3)

    slide_verificacion(prs, N())

    slide_panels(
        prs, N(), L('robustez', 'robustness'),
        L('Decisiones que evitan clases enteras de error',
          'Decisions that rule out whole classes of error'),
        None,
        [
         (L('Sin shell', 'No shell'),
          L('El lanzador se construye como lista de argumentos; shell=True no se usa nunca y los caracteres de control se rechazan.',
            'The launcher is built as an argument list; shell=True is never used and control characters are rejected.')),
         (L('Contenido', 'Contained'),
          L('El subproceso corre con cwd = jobs/ y las rutas de artefactos tienen que ser absolutas: no hay respaldo al directorio actual.',
            'The subprocess runs with cwd = jobs/ and artifact routes must be absolute: there is no fallback to the current directory.')),
         (L('Tres candados NOJOB', 'Three NOJOB guards'),
          L('Ni objetos de trabajo, ni archivos de solver, más una comprobación explícita de configuración.',
            'No job objects, no solver files, plus an explicit configuration check.')),
         (L('Detener significa detener', 'Cancel means cancel'),
          L('Se mata el árbol completo de procesos del solver y el proyecto queda marcado como CANCELLED, no como terminado.',
            'The whole solver process tree is killed and the project is marked CANCELLED, not finished.')),
         (L('Escrituras atómicas', 'Atomic writes'),
          L('Cada JSON se escribe en .tmp y se renombra, así que una caída no puede dejar medio manifiesto.',
            'Every JSON is written to .tmp and renamed, so a crash cannot leave half a manifest.')),
         (L('Nada con licencia se distribuye', 'Nothing licensed ships'),
          L('Los documentos DIN y FVA se quedan fuera del paquete por lista blanca, no por descuido.',
            'DIN and FVA documents stay out of the package by whitelist, not by accident.')),
        ],
        None,
        columns=3)

    slide_diagram(prs, N(), "dia_release.png", L("distribución", "distribution"))
    slide_diagram(prs, N(), "dia_estado.png", L("estado", "status"))
    slide_pasos(prs, N())
    slide_cierre(prs, N())

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    prs.save(OUTPUT)
    print("saved %s" % OUTPUT)
    print("slides: %d  language: %s" % (len(prs.slides._sldIdLst), LANG))
    return OUTPUT


if __name__ == "__main__":
    build()
